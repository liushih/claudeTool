from docx import Document
from pptx import Presentation

from runtime.render import parse_markdown, render_docx, render_pptx

MD = """# 周报

## 本周完成
- 完成提货单状态流转
- 修复登录超时

## 下周计划
- 接入用户管理
"""


def test_parse_markdown_splits_sections():
    title, intro, sections = parse_markdown(MD)
    assert title == "周报"
    assert intro == []
    assert [s[0] for s in sections] == ["本周完成", "下周计划"]
    assert sections[0][1] == ["完成提货单状态流转", "修复登录超时"]


PREAMBLE_MD = """# 周报

禅道数据缺失，本次仅基于代码提交。

## 本周完成
- 一
"""


def test_parse_markdown_keeps_preamble():
    _title, intro, _sections = parse_markdown(PREAMBLE_MD)
    assert intro == ["禅道数据缺失，本次仅基于代码提交。"]


def test_preamble_survives_rendering(tmp_path):
    docx_path = tmp_path / "p.docx"
    render_docx(PREAMBLE_MD, docx_path, "周报")
    docx_texts = [p.text for p in Document(docx_path).paragraphs]
    assert "禅道数据缺失，本次仅基于代码提交。" in docx_texts

    pptx_path = tmp_path / "p.pptx"
    render_pptx(PREAMBLE_MD, pptx_path, "周报", "sub")
    cover = "\n".join(
        shape.text for shape in Presentation(pptx_path).slides[0].shapes if shape.has_text_frame
    )
    assert "禅道数据缺失" in cover


def test_render_docx_writes_headings(tmp_path):
    out = tmp_path / "r.docx"
    render_docx(MD, out, "周报")
    assert out.exists() and out.stat().st_size > 0
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert "本周完成" in texts and "接入用户管理" in texts


def test_render_pptx_creates_slide_per_section(tmp_path):
    out = tmp_path / "r.pptx"
    render_pptx(MD, out, "周报", "hexincloud · 2026-W44")
    prs = Presentation(out)
    assert len(prs.slides) == 3
    first = "\n".join(shape.text for shape in prs.slides[1].shapes if shape.has_text_frame)
    assert "本周完成" in first


def test_render_cli_reports_missing_dependency(tmp_path, monkeypatch, capsys):
    import builtins

    from scripts import render as cli

    (tmp_path / "x.md").write_text(MD, encoding="utf-8")
    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name in {"docx", "pptx"}:
            raise ImportError(name)
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    rc = cli.main(["--input", str(tmp_path / "x.md"), "--format", "docx", "--out-dir", str(tmp_path)])
    assert rc == 1
    assert "缺少依赖" in capsys.readouterr().out
